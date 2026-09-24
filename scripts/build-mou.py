from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

GOLD = RGBColor(0xA9, 0x82, 0x20)
INK  = RGBColor(0x1A, 0x1A, 0x1C)
MUT  = RGBColor(0x5B, 0x5B, 0x5F)

doc = Document()

# --- page setup: A4, sensible margins -------------------------------------
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
for m in ("top_margin", "bottom_margin"): setattr(s, m, Cm(2.0))
for m in ("left_margin", "right_margin"): setattr(s, m, Cm(2.2))

# --- base styles ----------------------------------------------------------
n = doc.styles["Normal"]
n.font.name = "Calibri"; n.font.size = Pt(10.5); n.font.color.rgb = INK
n._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
pf = n.paragraph_format
pf.space_after = Pt(6); pf.line_spacing = 1.12

def style(name, size, bold=False, colour=INK, before=0, after=4, caps=False):
    st = doc.styles.add_style(name, 1)
    st.base_style = doc.styles["Normal"]
    st.font.size = Pt(size); st.font.bold = bold; st.font.color.rgb = colour
    st.font.name = "Calibri"
    if caps: st.font.all_caps = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    return st

style("GG Title",  20, True,  INK,  0, 2)
style("GG Sub",    10.5, False, MUT, 0, 14)
style("GG H1",     12, True,  GOLD, 14, 5, caps=True)
style("GG Party",  10.5, False, INK, 0, 8)
style("GG Foot",   8.5, False, MUT, 10, 0)

def p(text="", style_name=None, bold=False, size=None, colour=None, after=None, align=None):
    par = doc.add_paragraph(style=style_name)
    run = par.add_run(text)
    if bold: run.bold = True
    if size: run.font.size = Pt(size)
    if colour: run.font.color.rgb = colour
    if after is not None: par.paragraph_format.space_after = Pt(after)
    if align: par.alignment = align
    return par

def clause(num, text, bold_lead=None):
    """Numbered clause; optional bolded lead sentence."""
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(1.1)
    par.paragraph_format.first_line_indent = Cm(-1.1)
    par.paragraph_format.space_after = Pt(5)
    r = par.add_run(f"{num}\t"); r.bold = True
    if bold_lead:
        rb = par.add_run(bold_lead); rb.bold = True
        par.add_run(text)
    else:
        par.add_run(text)
    return par

def rule():
    par = doc.add_paragraph(); par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    pr = par._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:color"), "D6B65A")
    b.set(qn("w:space"), "1"); bd.append(b); pr.append(bd)

# ============================== HEADER =====================================
p("GET GOLD", "GG Title")
p("Vendor Pre-Agreement  ·  Memorandum of Understanding", "GG Sub")
rule()

# ============================== PARTIES ====================================
p("BETWEEN", "GG H1")
q = doc.add_paragraph(style="GG Party")
q.add_run("(1)  ").bold = True
q.add_run("Mohamed Amer").bold = True
q.add_run(", an individual resident in the United Arab Emirates, acting on his own behalf and on "
          "behalf of a company to be incorporated in the UAE to operate the ")
q.add_run("Get Gold").bold = True
q.add_run(" platform (“Get Gold”, “we”, “us”); and")

q = doc.add_paragraph(style="GG Party")
q.add_run("(2)  ").bold = True
q.add_run("______________________________________________")
q.add_run(", a business licensed in the UAE under trade licence number ")
q.add_run("____________________")
q.add_run(" (“the Vendor”, “you”).")

p("Date of signature:  ______ / ______ / 20______", after=10)
rule()

# ============================== CLAUSES ====================================
p("1.  Purpose and status", "GG H1")
clause("1.1", "Get Gold is building an online marketplace where customers in the UAE browse gold "
              "and jewellery from verified UAE vendors, see the live gold rate and every price "
              "component, reserve an item at a locked price, and either collect it or have it delivered.")
clause("1.2", "This document records what both parties intend. ",
        bold_lead=None)
par = doc.paragraphs[-1]
par.add_run("Except for clauses 8, 9, 10, 12 and 14, which are binding, this MOU is not legally "
            "binding").bold = True
par.add_run(" and neither party is obliged to transact.")
clause("1.3", "Nothing here creates a partnership, joint venture, employment or agency relationship, "
              "and neither party may hold itself out as able to bind the other.")

p("2.  Term", "GG H1")
clause("2.1", "This MOU takes effect on signature and expires on ")
doc.paragraphs[-1].add_run("30 April 2027").bold = True
doc.paragraphs[-1].add_run(", unless replaced earlier by a full agreement or ended under clause 11.")
clause("2.2", "On expiry, clauses 8, 9, 10, 12 and 14 survive for the periods stated in them.")

p("3.  What Get Gold provides", "GG H1")
clause("3.1", "Listing of your products on the Get Gold platform at no charge during Phase 1 (clause 5).")
clause("3.2", "Live gold pricing, so your prices update automatically with the market.")
clause("3.3", "Transparent display of each price component — gold value, making charge, certificate or "
              "assay fee, stone value, your premium, Get Gold service fee and delivery — shown separately to the customer.")
clause("3.4", "Customer reservations passed to you with the item, price and customer contact details.")
clause("3.5", "Product photography and listing setup for your initial stock, at no charge.")

p("4.  What the Vendor provides", "GG H1")
clause("4.1", "Accurate product information: weight, karat, making charge, stone value, certificate "
              "and hallmark details, and available quantity.")
clause("4.2", "", bold_lead="You honour the reserved price ")
doc.paragraphs[-1].add_run("shown to the customer. The initial stock and price request normally remains "
                           "open for ten (10) minutes. If you accept within that window, the displayed "
                           "total remains fixed until the payment deadline shown to the customer: initially "
                           "thirty (30) minutes for bank transfer and twenty-four (24) hours for cash or card "
                           "at delivery or collection.")
clause("4.3", "You keep stock levels current and tell us promptly when an item is no longer available.")
clause("4.4", "You respond to a reservation within the initial ten-minute window while accepting orders. "
              "An expired request cannot be revived without a new stock and price check.")
clause("4.5", "You hold a valid UAE trade licence permitting you to sell the goods listed, and all "
              "goods are authentic, accurately described, and hallmarked where required.")

p("5.  Commercial terms and customer service fee", "GG H1")
clause("5.1", "Get Gold charges you no listing fee, subscription fee or commission, including no commission "
              "on your making charge or Vendor Premium, for six (6) months from signature. No Vendor charge "
              "begins automatically when that period ends. Any future Vendor fee requires a separate written "
              "agreement accepted by both parties.", bold_lead="No Vendor commission during Phase 1.  ")
clause("5.2", "Get Gold may add a separately disclosed service fee paid by the customer. The Phase 1 standard "
              "rate is one percent (1%) of the merchandise subtotal, excluding delivery. Each new customer "
              "receives fifty percent (50%) off this service fee on their first three qualifying orders, producing "
              "an effective rate of one-half of one percent (0.5%). The exact percentage and amount must be shown "
              "before the order is placed.", bold_lead="Customer service fee.  ")
clause("5.3", "Where the customer pays you directly by cash, your card terminal or bank transfer, you collect "
              "the complete displayed order total. The separately identified Get Gold service-fee amount is "
              "collected by you on Get Gold's behalf and does not reduce the merchandise or delivery amount owed "
              "to you. You must not relabel, conceal or retain that service-fee amount as your sale proceeds.",
       bold_lead="Collection and remittance.  ")
clause("5.4", "Get Gold will provide a statement of completed orders and customer service fees collected. You "
              "will reconcile it against your invoices and payment records and remit undisputed amounts within "
              "________ days. Payment destination and tax-invoice requirements will be confirmed in writing after "
              "Get Gold is incorporated and before any remittance becomes due.")
clause("5.5", "Cancelled, rejected and expired unpaid orders carry no service fee. A full or partial customer "
              "refund produces the corresponding service-fee adjustment. A paid order keeps its introductory-order "
              "position even if later refunded; an unpaid cancelled, rejected or expired order does not.")
clause("5.6", "", bold_lead="No Get Gold fee is payable or remittable before Get Gold is incorporated in the UAE "
                            "and has notified you in writing that the customer-fee collection and settlement process is active.")

p("6.  Delivery", "GG H1")
clause("6.1", "Delivery is arranged and performed by you, at your election, by one of:")
for t in ["☐   (a)  Your own staff or vehicles, under your own insurance; or",
          "☐   (b)  An external delivery or courier company appointed by you, insured for the value of the goods."]:
    d = doc.add_paragraph(t); d.paragraph_format.left_indent = Cm(1.6); d.paragraph_format.space_after = Pt(3)
clause("6.2", "", bold_lead="The delivery charge to the customer must not exceed AED 60 (sixty UAE "
                            "dirhams) per order for deliveries within Dubai, ")
doc.paragraphs[-1].add_run("whichever method you choose. Get Gold displays this charge to the "
                           "customer and you collect it.")
clause("6.3", "Deliveries outside Dubai are by separate arrangement and are not covered by clause 6.2.")
clause("6.4", "Title and risk pass from you to the customer on delivery. Get Gold never takes "
              "possession, custody or title of any goods and is not a carrier, bailee or insurer.",
        bold_lead="The goods remain at your risk until delivered to the customer.  ")
clause("6.5", "You are responsible for insuring the goods in transit, whether carried by your own "
              "staff or by a company you appoint.")
clause("6.6", "The customer may choose store collection, which has no delivery fee, or an available delivery "
              "method. Payment choices may include cash, your card terminal and bank transfer to your verified "
              "business bank account. Online marketplace payment must not be offered until Get Gold confirms that "
              "an approved payment provider and settlement process are operational.")
clause("6.7", "For bank transfer, the customer waits for your stock acceptance, then sends the exact displayed "
              "amount and privately submits a transaction reference and receipt. A receipt image alone is not proof "
              "of payment. You must confirm cleared funds before the order proceeds. Late, short, duplicate or "
              "disputed transfers and related refunds remain your responsibility as seller.")

p("7.  Seller of record", "GG H1")
clause("7.1", "The contract of sale is between you and the customer. Get Gold introduces customers "
              "and coordinates; it does not buy, sell, hold or take title to gold.",
        bold_lead="You are the seller of record on every transaction.  ")
clause("7.2", "You are responsible for your own regulatory and AML obligations, including any "
              "obligations applicable to dealers in precious metals and stones in the UAE, and for "
              "any warranty, refund or consumer-protection obligation arising from your sales.")
clause("7.3", "You will indemnify Get Gold against any claim arising from the goods you sell, their "
              "description, authenticity, quality or delivery.")

p("8.  No circumvention  —  BINDING", "GG H1")
clause("8.1", "Where Get Gold introduces a customer to you, ")
doc.paragraphs[-1].add_run("you will not, for twelve (12) months from that introduction, "
                           "deliberately omit or misreport that customer’s transaction in order to avoid "
                           "an applicable Get Gold customer service fee or agreed platform process.").bold = True
clause("8.2", "This does not restrict you from serving customers who reach you independently, from "
              "your walk-in trade, or from your existing customer base.")
clause("8.3", "This clause is binding and survives expiry of this MOU.")

p("9.  Confidentiality  —  BINDING", "GG H1")
clause("9.1", "Each party will keep the other’s non-public commercial information confidential and "
              "use it only for the purposes of this MOU, for two (2) years from disclosure.")
clause("9.2", "This does not apply to information that is public, already known, or required to be "
              "disclosed by law or a regulator.")

p("10.  Intellectual property  —  BINDING", "GG H1")
clause("10.1", "Get Gold owns the platform, its software, brand, design, listing-page presentation, pricing "
               "displays and any photography Get Gold produces. Personal data remains subject to applicable "
               "privacy rights and is not transferred into Get Gold's ownership by this clause.")
clause("10.2", "You keep ownership of your own trademarks, product designs and any images you supply, "
               "and you grant Get Gold a non-exclusive, royalty-free licence to use your business "
               "name, logo and product images to list and market your products for the term of this MOU.")

p("11.  Ending this MOU, suspension and removal of listings", "GG H1")
clause("11.1", "", bold_lead="Get Gold may end this MOU at any time, with immediate effect, for any "
                             "reason or for no reason, ")
doc.paragraphs[-1].add_run("by written notice to you. Written notice includes email or WhatsApp to "
                           "the contact details in Schedule 1.")
clause("11.2", "You may end this MOU at any time, with immediate effect, by written notice to Get Gold.")
clause("11.3", "", bold_lead="Get Gold may at any time, without ending this MOU, and at its sole discretion:")
for t in ["(a)  remove, edit, suspend or decline to publish any listing;",
          "(b)  suspend your account temporarily; or",
          "(c)  decline to pass you any particular customer reservation."]:
    d = doc.add_paragraph(t); d.paragraph_format.left_indent = Cm(1.6); d.paragraph_format.space_after = Pt(3)
clause("11.4", "Get Gold will normally give reasons, but is not obliged to, and is not liable to you "
               "for any loss of sales, profit or opportunity arising from any action under clauses "
               "11.1 or 11.3.")
clause("11.5", "Without limiting the above, Get Gold may suspend you immediately if it reasonably "
               "believes that goods are not as described or not authentic, that a reserved price has "
               "not been honoured, that your trade licence has lapsed, or that any applicable law "
               "has been broken.", bold_lead="Immediate suspension for cause.  ")
clause("11.6", "Get Gold removes your listings; you honour any reservation a customer has already "
               "made and you have already confirmed, or refund the customer in full; and you settle "
               "any customer service fees properly collected on Get Gold's behalf and outstanding. Clauses 8, "
               "9, 10, 12 and 14 survive.",
        bold_lead="On termination:  ")
clause("11.7", "Neither party owes the other any compensation, penalty or payment merely for ending "
               "this MOU.")

p("12.  Assignment  —  BINDING", "GG H1")
clause("12.1", "", bold_lead="Get Gold may assign this MOU, and all rights and obligations under it, "
                             "to a UAE company incorporated by Mohamed Amer to operate the Get Gold "
                             "platform, without needing your further consent. ")
doc.paragraphs[-1].add_run("You agree that on such assignment that company replaces Mohamed Amer as a party.")
clause("12.2", "You may not assign this MOU without Get Gold’s written consent.")

p("13.  Non-exclusive", "GG H1")
clause("13.1", "This arrangement is non-exclusive on both sides. Get Gold may work with any number "
               "of other vendors, including your competitors. You may sell through any other channel, "
               "platform or premises.")

p("14.  Governing law  —  BINDING", "GG H1")
clause("14.1", "This MOU is governed by the laws of the United Arab Emirates as applied in the "
               "Emirate of Dubai, and the parties submit to the non-exclusive jurisdiction of the "
               "Dubai Courts.")
clause("14.2", "If this document is executed in English and Arabic and there is any conflict between "
               "them, the Arabic version prevails.")

# ============================== SCHEDULE ===================================
doc.add_page_break()
p("SCHEDULE 1  —  VENDOR DETAILS", "GG H1")
rows = ["Business name","Trade licence number","Licence expiry","Emirate","Store address",
        "Contact name","Mobile / WhatsApp","Email","VAT TRN (if registered)","Categories to list",
        "Approx. number of items at launch","Delivery method (clause 6.1)","Areas you will deliver to",
        "Payment methods offered","Bank-transfer beneficiary confirmed","Service-fee remittance period"]
t = doc.add_table(rows=len(rows), cols=2); t.style = "Table Grid"
for i, label in enumerate(rows):
    c0, c1 = t.rows[i].cells
    c0.text = label
    c0.paragraphs[0].runs[0].bold = True
    c0.paragraphs[0].runs[0].font.size = Pt(10)
    c0.width = Cm(6.4); c1.width = Cm(10.2)
    if label.startswith("Delivery method"):
        c1.text = "☐  Own staff          ☐  External company:  ______________________"
    elif label.startswith("Payment methods"):
        c1.text = "☐  Cash      ☐  Vendor card terminal      ☐  Bank transfer"
    elif label.startswith("Bank-transfer"):
        c1.text = "☐  Yes      ☐  Not offered"
    elif label.startswith("Service-fee"):
        c1.text = "________ days"
    t.rows[i].height = Cm(0.85)

p("", after=14)
p("SIGNED", "GG H1")

sg = doc.add_table(rows=1, cols=2); sg.autofit = True
left, right = sg.rows[0].cells
def sigblock(cell, who, sub):
    cell.paragraphs[0].text = ""
    r = cell.paragraphs[0].add_run(who); r.bold = True; r.font.size = Pt(10.5)
    a = cell.add_paragraph(sub); a.runs[0].font.size = Pt(8.5); a.runs[0].font.color.rgb = MUT
    for line in ["", "Signature:  ____________________", "Name:  ________________________",
                 "Position:  _____________________", "Date:  _________________________"]:
        q = cell.add_paragraph(line); q.paragraph_format.space_after = Pt(7)
        if q.runs: q.runs[0].font.size = Pt(10)
sigblock(left,  "For Get Gold", "Mohamed Amer, for himself and for a company to be incorporated")
sigblock(right, "For the Vendor", "Authorised signatory  ·  company stamp")

p("", after=6)
p("This Memorandum of Understanding is a pre-agreement. Except where a clause is marked BINDING, it "
  "does not create legally enforceable obligations. Both parties are advised to take independent "
  "legal advice before signing.", "GG Foot")

out = Path(__file__).resolve().parents[1] / "docs" / "Get-Gold-Vendor-Pre-Agreement.docx"
doc.save(out)
print("saved:", out)
