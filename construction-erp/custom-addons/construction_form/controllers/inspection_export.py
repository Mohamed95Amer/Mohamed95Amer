"""A completed inspection as a spreadsheet.

The PDF report has existed all along and is already in the print menu, which
is the right thing for issuing a signed record. It is the wrong thing for the
other half of what happens to a checklist: somebody wants the failed lines in
a column so they can sort them, count them, and paste them into whatever the
client asked for this week. That means a real spreadsheet, not a PDF and not
a CSV that Excel mangles the Arabic in.

Shaped after majal_documents/controllers/sheet_export.py, including the
guards, because a checklist is a worse place for them to be missing: an
inspector types free text into a comment box, and free text is where the
formula-injection payload arrives.
"""

import io
import re

from odoo import http
from odoo.http import content_disposition, request

# Anything a spreadsheet would treat as a formula rather than as text. Excel
# and LibreOffice both act on these, and one of them will prompt the person
# opening the file to run whatever the inspector typed.
FORMULA_LEADERS = ("=", "+", "-", "@", "\t", "\r")


def _safe(value):
    """Text that stays text when a spreadsheet opens it."""
    if value is None or value is False:
        return ""
    text = str(value)
    if text.startswith(FORMULA_LEADERS):
        return "'" + text
    return text


def _answer_text(answer):
    """One column for an answer, whatever kind of answer it is.

    A checklist mixes yes/no, numbers, dates, text and photographs, and a
    spreadsheet with five mostly-empty answer columns is unreadable. The
    typed value goes in one column and the header says which type it was.
    """
    kind = answer.answer_type
    if kind == "yes_no":
        return dict(
            answer._fields["answer_yes_no"]._description_selection(answer.env)
        ).get(answer.answer_yes_no, "")
    if kind == "number":
        return answer.answer_number
    if kind == "date":
        return answer.answer_date or ""
    if kind in ("photo", "signature"):
        # The image itself does not belong in a sheet meant for sorting and
        # counting; whether one was captured does.
        return "attached" if answer.answer_binary else ""
    return _safe(answer.answer_text)


class ConstructionInspectionExport(http.Controller):
    @http.route(
        "/construction/inspection/<int:inspection_id>/export.xlsx",
        type="http",
        auth="user",
        methods=["GET"],
    )
    def export_inspection(self, inspection_id):
        # search rather than browse: browse would hand back a record the
        # caller cannot read and only fail later, at field access, with an
        # error that says nothing about why.
        inspection = request.env["construction.form.inspection"].search(
            [("id", "=", inspection_id)], limit=1)
        if not inspection:
            return request.not_found()

        try:
            import openpyxl
        except ImportError:
            return request.make_response(
                "openpyxl is not installed on the server.",
                headers=[("Content-Type", "text/plain; charset=utf-8")],
                status=501,
            )

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Inspection"

        for label, value in (
            ("Inspection", inspection.name),
            ("Template", inspection.template_id.name),
            ("Project", inspection.project_id.display_name),
            ("Location", inspection.location),
            ("Inspector", inspection.inspector_id.display_name),
            ("Scheduled", inspection.scheduled_date or ""),
            ("Completed", inspection.completed_date or ""),
            ("Status", dict(
                inspection._fields["state"]._description_selection(inspection.env)
            ).get(inspection.state, "")),
            ("Score", inspection.score),
        ):
            sheet.append([label, _safe(value)])
        sheet.append([])

        header = sheet.max_row + 1
        sheet.append(["Section", "Question", "Type", "Answer", "Comment",
                      "Required", "Answered", "Failed"])
        for cell in sheet[header]:
            cell.font = openpyxl.styles.Font(bold=True)

        for answer in inspection.answer_ids.sorted(
                lambda a: (a.sequence, a.id)):
            sheet.append([
                _safe(answer.section),
                _safe(answer.question_id.name),
                _safe(answer.answer_type),
                _answer_text(answer),
                _safe(answer.comment),
                "Yes" if answer.is_required else "",
                "Yes" if answer.is_answered else "",
                "Yes" if answer.is_failed else "",
            ])

        # Freeze under the column headers so a fifty-question handover sheet
        # is still readable when scrolled.
        sheet.freeze_panes = sheet.cell(row=header + 1, column=1)
        for column, width in zip("ABCDEFGH", (22, 60, 12, 18, 40, 10, 10, 8)):
            sheet.column_dimensions[column].width = width

        stream = io.BytesIO()
        workbook.save(stream)
        filename = re.sub(r"[^A-Za-z0-9_.-]+", "-", inspection.name or "inspection")
        return request.make_response(
            stream.getvalue(),
            headers=[
                ("Content-Type", "application/vnd.openxmlformats-officedocument"
                                 ".spreadsheetml.sheet"),
                ("Content-Disposition",
                 content_disposition("%s.xlsx" % filename)),
                ("X-Content-Type-Options", "nosniff"),
                ("Cache-Control", "no-store"),
            ],
        )
